from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json

PATH='outputs/digital_asset_valuation_research_notebook.docx'
doc=Document(PATH)
BLUE='183A5A'; TEAL='167D7F'; PALE='F4F7F9'; WHITE='FFFFFF'; INK='17242E'; MUTED='5D6B75'; LIGHT='EAF1F5'
def shade(cell,fill):
 tc=cell._tc.get_or_add_tcPr(); s=OxmlElement('w:shd'); s.set(qn('w:fill'),fill); tc.append(s)
def margins(cell):
 tc=cell._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
 for tag,val in [('top',80),('start',105),('bottom',80),('end',105)]:
  e=OxmlElement('w:'+tag); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); mar.append(e)
 tc.append(mar)
def repeat(row):
 pr=row._tr.get_or_add_trPr(); e=OxmlElement('w:tblHeader'); e.set(qn('w:val'),'true'); pr.append(e)
def table(headers,rows,widths=None,font=8):
 t=doc.add_table(rows=1,cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False; repeat(t.rows[0])
 for j,h in enumerate(headers):
  c=t.rows[0].cells[j]; shade(c,BLUE); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(h); r.bold=True; r.font.color.rgb=RGBColor.from_string(WHITE); r.font.size=Pt(font)
 for i,row in enumerate(rows):
  cs=t.add_row().cells
  for j,v in enumerate(row):
   c=cs[j]; margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   if i%2: shade(c,PALE)
   p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.0; r=p.add_run(str(v)); r.font.size=Pt(font); r.font.color.rgb=RGBColor.from_string(INK)
 if widths:
  for row in t.rows:
   for j,w in enumerate(widths): row.cells[j].width=Inches(w)
 doc.add_paragraph()
def p(text='',style=None,bold=False,italic=False,color=None):
 q=doc.add_paragraph(style=style); r=q.add_run(text); r.bold=bold; r.italic=italic
 if color:r.font.color.rgb=RGBColor.from_string(color)
 return q
def bullets(xs):
 for x in xs:p(x,style='List Bullet')
def callout(label,text):
 t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,LIGHT); margins(c); q=c.paragraphs[0]; r=q.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL); q.add_run(text); doc.add_paragraph()

doc.add_page_break()
p('Notebook Update 1.1 — Pilot Universe Freeze',style='Heading 1')
p('Entry date: 22 August 2026 • Status: market snapshot frozen; pilot classifications populated; second-coder validation pending',italic=True,color=MUTED)
p('Selection decision',style='Heading 2')
p('The Phase 2 pilot universe is frozen at 25 non-stable cryptoassets and 16 stable-value assets. The selection combines market dominance, trading liquidity, public-data accessibility, economic-mechanism coverage, and a deliberately limited high-growth sleeve.')
callout('Coverage result','Using contemporaneous CoinGecko and DeFiLlama snapshots, the selected assets represent approximately 95.5% of total crypto market capitalization. The figure is indicative because vendor supply methodologies and snapshot timing differ; it should be recomputed from a synchronized production extract.')
p('Formal inclusion hierarchy',style='Heading 2')
bullets(['Tier 1 — Market core: large capitalization and meaningful trading volume; optimized for broad market coverage.', 'Tier 2 — Growth sleeve: expanding capitalization, supply, adoption, or economic relevance with accessible historical data.', 'Tier 3 — Economic coverage: adds a distinct value-accrual, consensus, privacy, governance, or stablecoin design mechanism.', 'Tier 4 — Failure control: retains failed/reflexive designs needed to estimate downside and avoid survivorship bias.', 'Exclusion rule: omit assets whose observed capitalization is primarily illiquid, whose volume is de minimis/unreliable, or whose historical/design data cannot support reproducible measurement.'])
p('Selected non-stable cryptoassets',style='Heading 2')
table(['Tier','Assets','Research role'],[
('Market core','BTC, ETH, BNB, XRP, SOL, TRX, DOGE, ZEC, LINK, ADA, XMR','Dominant monetary, infrastructure, payments, privacy, and oracle exposures'),
('Economic coverage','XLM, BCH, LTC, HBAR, AVAX, TON, UNI, AAVE, ARB, OP, POL','Additional consensus, L2, DeFi, governance, and supply-design variation'),
('Growth sleeve','HYPE, SUI, TAO','High-growth exchange/L1, scalable L1, and AI/service-network stories')],[1.25,2.4,3.2],8)
p('Selected stable-value assets',style='Heading 2')
table(['Tier','Assets','Research role'],[
('Market core','USDT, USDC, USDS, DAI, USDe, USD1, USDG','Captures the large majority of USD stablecoin supply across reserve, collateral, and synthetic structures'),
('Growth sleeve','RLUSD, PYUSD, USDD, GHO','Tests rapid issuance/adoption and newer distribution or protocol designs'),
('Structural coverage','FDUSD, TUSD, FRAX, PAXG','Adds issuer diversity, hybrid/protocol structure, and a commodity-referenced design'),
('Failure control','USTC / historical UST','Required reflexive-algorithmic failure observation')],[1.25,2.4,3.2],8)
p('Frozen snapshot metrics',style='Heading 2')
table(['Metric','Value','Interpretation'],[
('Total crypto market capitalization','$2.617 trillion','CoinGecko global snapshot'),('Selected non-stable crypto capitalization','$2.211 trillion','Approximately 84.5% of total market'),('Selected stable-value capitalization','$287.9 billion','DeFiLlama circulating peg-USD snapshot'),('Indicative combined coverage','Approximately 95.5%','Cross-provider estimate; synchronize for production'),('Workbook rows','25 crypto + 16 stable-value','41 asset-level pilot observations')],[2.25,1.65,2.95],8.2)
p('Deliverable and validation status',style='Heading 2')
bullets(['The editable workbook now contains Universe, asset_master, economic_design, stablecoin_design, Codebook, and Sources sheets.', 'Market fields are frozen to the retrieval date; formulas calculate liquidity ratios and value-accrual breadth.', 'Primary-source URLs are attached to design rows. These support the pilot coding but must be independently reviewed by a second coder.', 'Reserve-quality, redemption-friction, and transparency composite scores remain intentionally blank until component-level disclosures are extracted and scored.', 'No asset selection decision is currently required from the project sponsor. Adviser confirmation will be requested only if the desired sample size, historical start date, or treatment of memecoins/tokenized real-world assets changes materially.'])
p('Immediate next research action',style='Heading 2')
p('Conduct the second-coder review and fill the component-level stablecoin scoring inputs. In parallel, test data coverage for 2019–2026 daily market observations and the selected intraday depeg episodes. Assets failing the predeclared coverage threshold should be flagged—not silently replaced—before the final sample is locked.')
p('Update sources',style='Heading 2')
for s in ['CoinGecko Markets API: https://api.coingecko.com/api/v3/coins/markets','CoinGecko Global API: https://api.coingecko.com/api/v3/global','DeFiLlama Stablecoins API: https://stablecoins.llama.fi/stablecoins?includePrices=true']:
 p(s)

# version footer and core properties
for sec in doc.sections:
 for para in sec.footer.paragraphs:
  for run in para.runs:
   if 'Version 1.0' in run.text: run.text=run.text.replace('Version 1.0','Version 1.1')
doc.core_properties.comments='Version 1.1 adds the frozen Phase 2 pilot universe and selection methodology.'
doc.save(PATH)
print(PATH)
